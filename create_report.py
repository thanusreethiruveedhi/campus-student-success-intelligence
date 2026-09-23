from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ---------------------------------------------------------
# CREATE DOCUMENT
# ---------------------------------------------------------

doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# ---------------------------------------------------------
# STYLES
# ---------------------------------------------------------

styles = doc.styles

styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True

styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].font.bold = True

styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True


# ---------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------

def add_paragraph(text="", bold=False, italic=False, align=None):
    p = doc.add_paragraph()

    if align:
        p.alignment = align

    run = p.add_run(text)
    run.bold = bold
    run.italic = italic

    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_table(headers, rows):
    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Rows
    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_paragraph()

    return table


def add_page_number():
    section = doc.sections[0]
    footer = section.footer

    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("Page ")

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


# ---------------------------------------------------------
# TITLE PAGE
# ---------------------------------------------------------

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = p.add_run("\n\n\n")
run.font.size = Pt(12)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "CAMPUS RESOURCE & STUDENT SUCCESS\n"
    "INTELLIGENCE SYSTEM"
)
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(24)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "\nBusiness Intelligence & Predictive Analytics Project"
)
run.bold = True
run.font.size = Pt(15)

doc.add_paragraph("\n")

add_paragraph(
    "Project Report",
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER
)

doc.add_paragraph("\n")

add_paragraph(
    "Submitted by",
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_paragraph(
    "Thanusree",
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_paragraph(
    "B.Tech – CSE with AI Specialization",
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_paragraph(
    "Ashoka Women's Engineering College",
    align=WD_ALIGN_PARAGRAPH.CENTER
)

doc.add_paragraph("\n")

add_paragraph(
    "Academic Project – 2026",
    align=WD_ALIGN_PARAGRAPH.CENTER
)

doc.add_page_break()


# ---------------------------------------------------------
# ABSTRACT
# ---------------------------------------------------------

doc.add_heading("1. Abstract", level=1)

add_paragraph(
    "The Campus Resource & Student Success Intelligence System is a "
    "Business Intelligence and Predictive Analytics project designed to "
    "analyze student academic performance, attendance, learning engagement, "
    "campus resource usage, and academic risk."
)

add_paragraph(
    "The project uses a synthetic dataset containing 1,000 student records "
    "and 21 variables. Exploratory Data Analysis techniques are used to "
    "identify patterns in attendance, previous CGPA, internal marks, "
    "assignment performance, LMS activity, study hours, backlogs, and "
    "student-development activities."
)

add_paragraph(
    "A Random Forest classification model is developed to classify students "
    "into Low Risk, Medium Risk, and High Risk categories. The model achieved "
    "88.50% accuracy and 84.36% macro F1-score on a 200-student test set."
)

add_paragraph(
    "The project combines data preparation, exploratory analysis, business "
    "insights, predictive analytics, machine learning, and interactive "
    "dashboard concepts into a single student-success intelligence solution."
)



# ---------------------------------------------------------
# INTRODUCTION
# ---------------------------------------------------------

doc.add_heading("2. Introduction", level=1)

add_paragraph(
    "Educational institutions generate large amounts of student-related "
    "information through academic assessments, attendance systems, learning "
    "management systems, libraries, laboratories, projects, certifications, "
    "internships, and other student activities."
)

add_paragraph(
    "Business Intelligence can transform such raw information into useful "
    "insights that support academic monitoring and institutional decision "
    "making. Predictive Analytics can further help identify patterns that "
    "may be associated with academic risk."
)

add_paragraph(
    "This project applies a complete analytics workflow beginning with data "
    "preparation and exploratory analysis and continuing through business "
    "insight generation, predictive modelling, model evaluation, and "
    "dashboard-oriented visualization."
)


# ---------------------------------------------------------
# PROBLEM STATEMENT
# ---------------------------------------------------------

doc.add_heading("3. Problem Statement", level=1)

add_paragraph(
    "Institutions need structured ways to understand academic performance "
    "and student engagement across multiple dimensions. Looking at academic "
    "marks alone may not provide a complete view of student activity."
)

add_paragraph(
    "The problem addressed by this project is to build an analytical system "
    "that combines academic, attendance, engagement, resource-usage, and "
    "student-development variables to identify meaningful patterns and "
    "develop a prototype for academic-risk classification."
)


# ---------------------------------------------------------
# OBJECTIVES
# ---------------------------------------------------------

doc.add_heading("4. Objectives", level=1)

objectives = [
    "Analyze student academic performance using descriptive statistics and visualization.",
    "Examine the association between attendance and Final CGPA.",
    "Analyze the relationship between backlogs and academic performance.",
    "Study LMS activity and other engagement indicators.",
    "Analyze campus resource usage and student-development activities.",
    "Identify characteristics of students classified as High Risk.",
    "Compare academic risk across different branches.",
    "Build a Random Forest model for academic-risk classification.",
    "Evaluate the predictive model using accuracy, F1-score, and confusion matrix.",
    "Present the results in an interactive Business Intelligence dashboard."
]

for item in objectives:
    add_bullet(item)


# ---------------------------------------------------------
# BUSINESS QUESTIONS
# ---------------------------------------------------------

doc.add_heading("5. Business Questions", level=1)

questions = [
    "What is the average Final CGPA of students?",
    "How does attendance relate to Final CGPA?",
    "How does previous CGPA relate to Final CGPA?",
    "How are backlogs associated with academic performance?",
    "How is LMS activity associated with Final CGPA?",
    "Which branches have different levels of academic performance?",
    "How many students fall into each academic risk category?",
    "What characteristics are observed among High-Risk students?",
    "How does High-Risk student percentage vary by branch?",
    "Which variables are most important to the predictive model?"
]

for item in questions:
    add_bullet(item)


# ---------------------------------------------------------
# DATASET
# ---------------------------------------------------------

doc.add_heading("6. Dataset Description", level=1)

add_paragraph(
    "The project uses a synthetic student dataset containing 1,000 records "
    "and 21 columns. Each row represents a student record."
)

dataset_rows = [
    ["Student_ID", "Unique student identifier"],
    ["Branch", "Academic branch"],
    ["Year", "Academic year"],
    ["Gender", "Student gender category"],
    ["Attendance", "Attendance percentage"],
    ["Previous_CGPA", "Previous academic CGPA"],
    ["Internal_Marks", "Internal assessment marks"],
    ["Assignment_Score", "Assignment performance score"],
    ["Study_Hours", "Average study hours"],
    ["LMS_Activity", "Learning Management System activity"],
    ["Library_Visits", "Number of library visits"],
    ["Lab_Usage", "Laboratory usage indicator"],
    ["Certifications", "Number of certifications"],
    ["Projects", "Number of projects"],
    ["Hackathons", "Number of hackathons"],
    ["Internships", "Number of internships"],
    ["Commute_Time", "Approximate commute time"],
    ["Backlogs", "Number of backlogs"],
    ["Final_CGPA", "Final academic CGPA"],
    ["Academic_Status", "Academic performance status"],
    ["Risk_Level", "Low, Medium, or High academic risk"]
]

add_table(
    ["Variable", "Description"],
    dataset_rows
)


# ---------------------------------------------------------
# DATA QUALITY
# ---------------------------------------------------------

doc.add_heading("7. Data Preparation and Quality Check", level=1)

add_paragraph(
    "Before analysis, the dataset was checked for missing values, duplicate "
    "records, duplicate student identifiers, and reasonable numerical ranges."
)

quality_rows = [
    ["Total Records", "1,000"],
    ["Total Variables", "21"],
    ["Missing Values", "0"],
    ["Duplicate Rows", "0"],
    ["Duplicate Student IDs", "0"],
    ["Attendance Range", "50.0% – 100.0%"],
    ["Previous CGPA Range", "5.0 – 9.8"],
    ["Final CGPA Range", "5.26 – 9.05"]
]

add_table(
    ["Data Quality Check", "Result"],
    quality_rows
)


# ---------------------------------------------------------
# EDA
# ---------------------------------------------------------

doc.add_heading("8. Exploratory Data Analysis", level=1)

add_paragraph(
    "Exploratory Data Analysis was performed using Python, Pandas, "
    "Matplotlib, and Seaborn. The analysis examined academic status, "
    "risk distribution, branch performance, attendance, backlogs, LMS "
    "activity, resource usage, and correlations."
)

doc.add_heading("8.1 Overall Academic Performance", level=2)

add_table(
    ["Metric", "Value"],
    [
        ["Total Students", "1,000"],
        ["Average Final CGPA", "7.33"],
        ["Average Attendance", "80.25%"],
        ["Average Previous CGPA", "7.51"],
        ["Average Backlogs", "0.38"]
    ]
)

doc.add_heading("8.2 Academic Status Distribution", level=2)

add_table(
    ["Academic Status", "Students"],
    [
        ["Good Standing", "422"],
        ["Needs Attention", "555"],
        ["At Risk", "23"]
    ]
)

doc.add_heading("8.3 Risk Distribution", level=2)

add_table(
    ["Risk Level", "Students"],
    [
        ["Low Risk", "650"],
        ["Medium Risk", "230"],
        ["High Risk", "120"]
    ]
)

doc.add_heading("8.4 Branch Performance", level=2)

add_table(
    ["Branch", "Students", "Average CGPA", "Average Attendance"],
    [
        ["Mechanical", "172", "7.42", "80.99%"],
        ["EEE", "174", "7.33", "80.07%"],
        ["Civil", "155", "7.33", "80.15%"],
        ["ECE", "154", "7.31", "79.14%"],
        ["CSE", "181", "7.30", "80.78%"],
        ["CSE-AI", "164", "7.27", "80.21%"]
    ]
)

doc.add_heading("8.5 Attendance Analysis", level=2)

add_table(
    ["Attendance Group", "Students", "Average CGPA"],
    [
        ["Below 60%", "17", "6.64"],
        ["60–75%", "282", "7.15"],
        ["75–85%", "393", "7.32"],
        ["Above 85%", "308", "7.53"]
    ]
)

doc.add_heading("8.6 Backlog Analysis", level=2)

add_table(
    ["Backlogs", "Students", "Average CGPA"],
    [
        ["0", "687", "7.49"],
        ["1", "248", "7.04"],
        ["2", "61", "6.71"],
        ["3", "3", "6.57"],
        ["4", "1", "6.81"]
    ]
)


# ---------------------------------------------------------
# BUSINESS INSIGHTS
# ---------------------------------------------------------

doc.add_heading("9. Business Insights", level=1)

doc.add_heading("9.1 Attendance and Academic Performance", level=2)

add_paragraph(
    "Students with attendance above 85% have the highest average CGPA "
    "(7.53), while students below 60% have the lowest average CGPA (6.64). "
    "The correlation between Attendance and Final CGPA is 0.276."
)

add_paragraph(
    "These results indicate an association in the synthetic dataset and "
    "should not be interpreted as evidence that attendance directly causes "
    "higher academic performance."
)

doc.add_heading("9.2 Backlogs and Academic Performance", level=2)

add_paragraph(
    "Students with no backlogs have an average CGPA of 7.49, compared with "
    "7.04 for students with one backlog and 6.71 for students with two "
    "backlogs. The correlation between Backlogs and Final CGPA is -0.401."
)

doc.add_heading("9.3 LMS Engagement", level=2)

add_paragraph(
    "LMS activity has a positive association with Final CGPA in the dataset. "
    "The correlation between LMS Activity and Final CGPA is 0.260."
)

doc.add_heading("9.4 High-Risk Student Profile", level=2)

add_table(
    ["Metric", "High-Risk Students", "Overall"],
    [
        ["Students", "120", "1,000"],
        ["Average Attendance", "76.79%", "80.25%"],
        ["Average Previous CGPA", "6.34", "7.51"],
        ["Average Final CGPA", "6.49", "7.33"],
        ["Average Internal Marks", "69.20", "71.69"],
        ["Average Assignment Score", "72.91", "74.20"],
        ["Average LMS Activity", "64.54", "67.92"],
        ["Average Study Hours", "4.15", "4.54"],
        ["Average Backlogs", "1.46", "0.38"]
    ]
)

doc.add_heading("9.5 High-Risk Students by Branch", level=2)

add_table(
    ["Branch", "Students", "High-Risk Students", "High-Risk %"],
    [
        ["CSE-AI", "164", "24", "14.63%"],
        ["ECE", "154", "21", "13.64%"],
        ["EEE", "174", "22", "12.64%"],
        ["Civil", "155", "18", "11.61%"],
        ["CSE", "181", "19", "10.50%"],
        ["Mechanical", "172", "16", "9.30%"]
    ]
)


# ---------------------------------------------------------
# PREDICTIVE ANALYTICS
# ---------------------------------------------------------

doc.add_heading("10. Predictive Analytics and Machine Learning", level=1)

add_paragraph(
    "A Random Forest Classifier was developed to predict the Risk_Level "
    "of students. The target contains three categories: Low Risk, Medium "
    "Risk, and High Risk."
)

add_paragraph(
    "Fourteen input features were used. Final_CGPA and Academic_Status "
    "were excluded from the model features to avoid target leakage."
)

features_text = (
    "Attendance, Previous_CGPA, Internal_Marks, Assignment_Score, "
    "Study_Hours, LMS_Activity, Library_Visits, Lab_Usage, Certifications, "
    "Projects, Hackathons, Internships, Commute_Time, and Backlogs."
)

add_paragraph(
    "Model features: " + features_text
)

doc.add_heading("10.1 Train-Test Split", level=2)

add_table(
    ["Dataset", "Students"],
    [
        ["Training Set", "800"],
        ["Testing Set", "200"]
    ]
)

doc.add_heading("10.2 Random Forest Configuration", level=2)

add_table(
    ["Parameter", "Value"],
    [
        ["Algorithm", "Random Forest Classifier"],
        ["Number of Trees", "200"],
        ["Maximum Depth", "10"],
        ["Minimum Samples Split", "5"],
        ["Class Weight", "Balanced"],
        ["Random State", "42"]
    ]
)


# ---------------------------------------------------------
# MODEL EVALUATION
# ---------------------------------------------------------

doc.add_heading("11. Model Evaluation", level=1)

add_table(
    ["Metric", "Result"],
    [
        ["Accuracy", "88.50%"],
        ["Macro F1", "84.36%"],
        ["Test Students", "200"]
    ]
)

doc.add_heading("11.1 Classification Report", level=2)

add_table(
    ["Risk Level", "Precision", "Recall", "F1-Score", "Support"],
    [
        ["High Risk", "0.778", "0.875", "0.824", "24"],
        ["Low Risk", "0.975", "0.915", "0.944", "130"],
        ["Medium Risk", "0.725", "0.804", "0.763", "46"]
    ]
)

doc.add_heading("11.2 Confusion Matrix", level=2)

add_table(
    ["Actual / Predicted", "Low Risk", "Medium Risk", "High Risk"],
    [
        ["Low Risk", "119", "11", "0"],
        ["Medium Risk", "3", "37", "6"],
        ["High Risk", "0", "3", "21"]
    ]
)

doc.add_heading("11.3 Feature Importance", level=2)

add_table(
    ["Feature", "Importance"],
    [
        ["Previous_CGPA", "0.2879"],
        ["Backlogs", "0.2871"],
        ["Attendance", "0.0694"],
        ["Study_Hours", "0.0638"],
        ["Internal_Marks", "0.0513"],
        ["LMS_Activity", "0.0451"],
        ["Commute_Time", "0.0378"],
        ["Assignment_Score", "0.0354"],
        ["Lab_Usage", "0.0351"],
        ["Library_Visits", "0.0207"],
        ["Projects", "0.0199"],
        ["Hackathons", "0.0173"],
        ["Certifications", "0.0157"],
        ["Internships", "0.0136"]
    ]
)

add_paragraph(
    "Previous_CGPA and Backlogs were the two features with the highest "
    "model importance in this synthetic dataset."
)


# ---------------------------------------------------------
# DASHBOARD
# ---------------------------------------------------------

doc.add_heading("12. Interactive Dashboard", level=1)

add_paragraph(
    "A Streamlit dashboard was developed to provide an interactive view "
    "of the project's Business Intelligence findings."
)

add_paragraph(
    "The dashboard contains four main sections:"
)

dashboard_sections = [
    "Executive Overview – key performance indicators, academic status, risk distribution, and branch performance.",
    "Academic Intelligence – CGPA, attendance, backlogs, branch performance, and related visualizations.",
    "Student Engagement – LMS activity, campus resources, study hours, and engagement correlations.",
    "Risk & Prediction – model performance, confusion matrix, feature importance, branch risk analysis, and individual student prediction."
]

for item in dashboard_sections:
    add_bullet(item)


# ---------------------------------------------------------
# TECHNOLOGIES
# ---------------------------------------------------------

doc.add_heading("13. Technologies Used", level=1)

technologies = [
    ["Python", "Programming and data analysis"],
    ["Pandas", "Data manipulation and analysis"],
    ["NumPy", "Numerical computation"],
    ["Matplotlib", "Data visualization"],
    ["Seaborn", "Statistical visualization"],
    ["Plotly", "Interactive visualization"],
    ["Scikit-learn", "Machine learning"],
    ["Streamlit", "Interactive dashboard"],
    ["Jupyter Notebook", "Project documentation and analysis"],
    ["GitHub Codespaces", "Cloud development environment"],
    ["Git & GitHub", "Version control and project hosting"]
]

add_table(
    ["Technology", "Purpose"],
    technologies
)


# ---------------------------------------------------------
# PROJECT WORKFLOW
# ---------------------------------------------------------

doc.add_heading("14. Project Workflow", level=1)

workflow = [
    "Raw Data Collection",
    "Data Cleaning and Quality Check",
    "Business Questions",
    "Exploratory Data Analysis",
    "Business Insights",
    "Predictive Modelling",
    "Model Evaluation",
    "Interactive Dashboard",
    "Decision-Support Insights"
]

for i, step in enumerate(workflow, start=1):
    add_number(step)


# ---------------------------------------------------------
# LIMITATIONS
# ---------------------------------------------------------

doc.add_heading("15. Limitations", level=1)

limitations = [
    "The dataset is synthetic and may not represent real student populations.",
    "The observed relationships are associations rather than causal relationships.",
    "The Risk_Level target was generated using variables that are also used as model features.",
    "Therefore, the model performance should be considered a prototype result rather than evidence of real-world predictive performance.",
    "Only one train-test split was used for evaluation.",
    "Actual institutional deployment would require privacy, governance, fairness, validation, and monitoring procedures."
]

for item in limitations:
    add_bullet(item)


# ---------------------------------------------------------
# FUTURE SCOPE
# ---------------------------------------------------------

doc.add_heading("16. Future Scope", level=1)

future_scope = [
    "Integrate appropriately anonymized real institutional data.",
    "Develop time-based early-warning monitoring.",
    "Add explainable AI methods for individual predictions.",
    "Track the outcomes of academic interventions.",
    "Perform fairness and model-monitoring analysis.",
    "Add automated academic-support alerts.",
    "Deploy the Streamlit dashboard on a cloud platform.",
    "Expand the system with additional institutional resource data."
]

for item in future_scope:
    add_bullet(item)


# ---------------------------------------------------------
# CONCLUSION
# ---------------------------------------------------------

doc.add_heading("17. Conclusion", level=1)

add_paragraph(
    "The Campus Resource & Student Success Intelligence System demonstrates "
    "how Business Intelligence and Predictive Analytics can be combined "
    "to analyze multiple dimensions of student performance and engagement."
)

add_paragraph(
    "The exploratory analysis identified meaningful associations involving "
    "attendance, LMS activity, backlogs, previous CGPA, and academic "
    "performance. The Random Forest prototype achieved 88.50% accuracy "
    "and 84.36% macro F1 on the synthetic test dataset."
)

add_paragraph(
    "The project provides an end-to-end analytical workflow covering data "
    "preparation, exploratory analysis, insight generation, machine "
    "learning, evaluation, and dashboard visualization."
)

add_paragraph(
    "Because the project uses synthetic data and a constructed risk target, "
    "the results should be treated as a demonstration of the analytics "
    "workflow rather than as a validated real-world student-risk system."
)


# ---------------------------------------------------------
# REFERENCES
# ---------------------------------------------------------

doc.add_heading("18. References", level=1)

references = [
    "Python Documentation – https://www.python.org/",
    "Pandas Documentation – https://pandas.pydata.org/",
    "NumPy Documentation – https://numpy.org/",
    "Scikit-learn Documentation – https://scikit-learn.org/",
    "Matplotlib Documentation – https://matplotlib.org/",
    "Seaborn Documentation – https://seaborn.pydata.org/",
    "Plotly Documentation – https://plotly.com/python/",
    "Streamlit Documentation – https://streamlit.io/"
]

for reference in references:
    add_bullet(reference)


# ---------------------------------------------------------
# FOOTER
# ---------------------------------------------------------

add_page_number()


# ---------------------------------------------------------
# SAVE DOCUMENT
# ---------------------------------------------------------

output_file = "Thanusree_ProjectReport.docx"

doc.save(output_file)

print("=" * 60)
print("PROJECT REPORT CREATED SUCCESSFULLY")
print("=" * 60)
print(f"File: {output_file}")
print("Format: Microsoft Word (.docx)")
print("=" * 60)